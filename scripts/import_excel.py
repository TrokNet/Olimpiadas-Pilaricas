"""Carga transaccional Excel -> SQLite/PostgreSQL. Nunca escribe en el Excel."""
import argparse
import hashlib
import json
import re
import sys
import unicodedata
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from uuid import uuid5, NAMESPACE_URL
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from sqlalchemy import delete
from sqlalchemy.orm import Session
from backend.database import ROOT, Grade, Entry, Match, Metadata, engine_for, initialize

SPORTS = {'Fútbol 6':'futbol-masculino','Vóley masculino':'voley-masculino','Vóley femenino':'voley-femenino','Básquet masculino':'basquet-masculino','Básquet femenino':'basquet-femenino','Básquet (reserva)':'basquet-reserva'}
ORDER = ['3 Años','4 Años','5 Años','Primero Primaria','Segundo Primaria','Tercero Primaria','Cuarto Primaria','Quinto Primaria','Sexto Primaria','Primero Secundaria','Segundo Secundaria','Tercero Secundaria','Cuarto Secundaria','Quinto Secundaria','Ex-Alumnas']

def slug(value):
    value = unicodedata.normalize('NFKD', value).encode('ascii','ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+','-',value).strip('-')

def get_rows(sheet, header_row):
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < header_row: raise ValueError(f'{sheet.title}: faltan encabezados')
    headers = [str(v).strip() if v is not None else '' for v in rows[header_row-1]]
    for index, values in enumerate(rows[header_row:], header_row+1):
        if any(v is not None for v in values): yield index, dict(zip(headers, values))

def date_value(value):
    if isinstance(value, datetime): return value.date().isoformat()
    if isinstance(value, date): return value.isoformat()
    if isinstance(value, str): return date.fromisoformat(value.strip()).isoformat()
    raise ValueError(f'Fecha no válida: {value}')

def time_value(value):
    if isinstance(value, (datetime, time)): return value.strftime('%H:%M')
    if isinstance(value, (float,int)): return (datetime(2000,1,1)+timedelta(days=value)).strftime('%H:%M')
    if isinstance(value,str): return datetime.strptime(value.strip(),'%H:%M').strftime('%H:%M')
    raise ValueError(f'Hora no válida: {value}')

def score(value):
    if value is None or value == '': return None
    n=int(value)
    if n < 0 or float(value)!=n: raise ValueError('Marcador debe ser entero no negativo')
    return n

def import_excel(path, engine, config=None):
    config=config or json.loads((ROOT/'data/import_config.json').read_text(encoding='utf-8-sig'))
    wb=load_workbook(path,read_only=True,data_only=True)
    c=config['columns']; hr=config['header_row']
    # Correcciones aprobadas posteriores al PDF/Excel, separadas por identidad.
    name_overrides=json.loads((ROOT/'data/grade_name_overrides.json').read_text(encoding='utf-8'))
    entries={}; grades={}; matches=[]; errors=[]; keys=set()
    for row, values in get_rows(wb[config['teams_sheet']],hr):
        try:
            for field in ['sport','number','grade','name','color']:
                if values.get(c[field]) is None: raise ValueError(f'Falta {c[field]}')
            sport=SPORTS[str(values[c['sport']])]; n=int(values[c['number']]); grade=str(values[c['grade']]).strip(); gid=slug(grade)
            if n<1: raise ValueError('TEAM debe ser positivo')
            eid=f'{sport}:TEAM-{n}'
            if eid in entries or any(e['sport']==sport and e['grade_id']==gid for e in entries.values()): raise ValueError('TEAM o grado duplicado en disciplina')
            g={'id':gid,'label':grade,'name':str(values[c['name']]).strip(),'color':str(values[c['color']]).strip(),'display_order':ORDER.index(grade) if grade in ORDER else 99}
            g['name']=name_overrides.get(gid,g['name'])
            if gid in grades and grades[gid]!=g: raise ValueError('Nombre/color inconsistente para el grado')
            grades[gid]=g
            entries[eid]={'id':eid,'sport':sport,'team_number':n,'grade_id':gid,'source_row':row}
        except (ValueError,KeyError,TypeError) as e: errors.append(f'Equipos fila {row}: {e}')
    for row, values in get_rows(wb[config['matches_sheet']],hr):
        try:
            for field in ['date','court','start','end','sport','phase','home','away','kind']:
                if values.get(c[field]) is None: raise ValueError(f'Falta {c[field]}')
            sport=SPORTS[str(values[c['sport']])]; d=date_value(values[c['date']]); start=time_value(values[c['start']]); end=time_value(values[c['end']])
            if start>=end: raise ValueError('Inicio debe ser anterior al fin')
            labels=[str(values[c[k]]).strip() for k in ['home','away']]; ids=[]
            for label in labels:
                m=re.fullmatch(r'(?:Equipo|TEAM)\s+(\d+)',label,re.I)
                eid=f'{sport}:TEAM-{int(m[1])}' if m else None
                if eid and eid not in entries: raise ValueError(f'{label} sin correspondencia en {sport}')
                ids.append(eid)
            if ids[0] and ids[0]==ids[1]: raise ValueError('Un equipo no puede jugar contra sí mismo')
            phase=str(values[c['phase']]); kind=str(values[c['kind']])
            identity='|'.join([sport,phase,*sorted(labels)])
            if kind=='RESERVA': identity+=f'|{d}|{start}|{values[c["court"]]}'
            if identity in keys: raise ValueError('Partido duplicado')
            keys.add(identity)
            hs=score(values.get(c['home_score'])); aws=score(values.get(c['away_score']))
            if (hs is None)!=(aws is None): raise ValueError('Ambos marcadores son obligatorios si se registra un resultado')
            matches.append({'id':str(uuid5(NAMESPACE_URL,'nsp-2026|'+identity)),'date':d,'start':start,'end':end,'sport':sport,'phase':phase,'court':str(values[c['court']]),'home_id':ids[0],'away_id':ids[1],'home_label':labels[0],'away_label':labels[1],'kind':kind,'referee':str(values.get(c['referee']) or 'Por asignar'),'home_score':hs,'away_score':aws,'source_row':row,'source_sheet':config['matches_sheet'],'raw_json':json.dumps(values,ensure_ascii=False,default=str)})
        except (ValueError,KeyError,TypeError) as e: errors.append(f'Fixture fila {row}: {e}')
    wb.close()
    if not entries or not matches: errors.append('No se puede importar una fuente sin equipos o partidos')
    if errors: raise ValueError('\n'.join(errors))
    initialize(engine)
    with Session(engine) as session, session.begin():
        # Una única transacción conserva la versión previa ante cualquier error.
        session.execute(delete(Match)); session.execute(delete(Entry)); session.execute(delete(Grade))
        session.add_all(Grade(**g) for g in grades.values()); session.flush()
        session.add_all(Entry(**e) for e in entries.values()); session.flush()
        session.add_all(Match(**m) for m in matches)
        for key,value in {'schema_version':'1','imported_at':datetime.now(timezone.utc).isoformat(),'source_file':Path(path).name,'source_sha256':hashlib.sha256(Path(path).read_bytes()).hexdigest()}.items(): session.merge(Metadata(key=key,value=value))
    return {'grades':len(grades),'entries':len(entries),'matches':len(matches)}

if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--excel',default=str(ROOT/'data/olimpiadas.xlsx')); parser.add_argument('--database-url'); args=parser.parse_args()
    try: print(json.dumps(import_excel(args.excel,engine_for(args.database_url)),ensure_ascii=False))
    except Exception as e: print(f'IMPORTACIÓN RECHAZADA:\n{e}',file=sys.stderr); sys.exit(1)
